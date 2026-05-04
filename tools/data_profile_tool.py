from __future__ import annotations
import csv
import json
import math
import statistics
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from tools.base import BaseTool, ToolInput, ToolResult

def _to_float(v: Any) -> Optional[float]:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None
        return float(v)
    text = str(v).strip()
    if not text:
        return None
    try:
        return float(text.replace(",", ""))
    except Exception:
        return None

def _read_csv(path: Path) -> Tuple[List[str], List[Dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        columns = reader.fieldnames or []
        rows = [dict(r) for r in reader]
    return list(columns), rows

def _read_json(path: Path) -> Tuple[List[str], List[Dict[str, Any]], Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        rows = [r for r in data if isinstance(r, dict)]
        cols = sorted({k for r in rows for k in r.keys()})
        return cols, rows, data
    if isinstance(data, dict):
        # 支持 {"data":[...]} 常见结构
        maybe = data.get("data")
        if isinstance(maybe, list) and all(isinstance(i, dict) for i in maybe):
            cols = sorted({k for r in maybe for k in r.keys()})
            return cols, list(maybe), data
        return [], [], data
    return [], [], data

def _read_txt(path: Path) -> Dict[str, Any]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    lines = text.splitlines()
    words = [w.strip(".,;:!?()[]{}\"'").lower() for w in text.split()]
    words = [w for w in words if len(w) > 1]
    top_words = Counter(words).most_common(10)
    return {
        "type": "text",
        "char_count": len(text),
        "line_count": len(lines),
        "preview_lines": lines[:12],
        "top_words": top_words,
    }

def _read_xlsx(path: Path, sheet_name: Optional[str] = None) -> Tuple[List[str], List[Dict[str, Any]], str]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("读取 Excel 需要 openpyxl 依赖，请先安装 openpyxl。") from exc

    wb = load_workbook(filename=str(path), data_only=True, read_only=True)
    target = wb[sheet_name] if (sheet_name and sheet_name in wb.sheetnames) else wb[wb.sheetnames[0]]
    rows = list(target.iter_rows(values_only=True))
    if not rows:
        return [], [], target.title
    headers = [str(h).strip() if h is not None else f"col_{idx + 1}" for idx, h in enumerate(rows[0])]
    data_rows: List[Dict[str, Any]] = []
    for row in rows[1:]:
        row_map = {}
        for idx, col in enumerate(headers):
            row_map[col] = row[idx] if idx < len(row) else None
        data_rows.append(row_map)
    return headers, data_rows, target.title

def _profile_table(columns: List[str], rows: List[Dict[str, Any]]) -> Dict[str, Any]:
    total = len(rows)
    col_stats: List[Dict[str, Any]] = []
    warnings: List[str] = []
    numeric_summary: Dict[str, Dict[str, Any]] = {}
    category_summary: Dict[str, List[Tuple[str, int]]] = {}

    for col in columns:
        values = [r.get(col) for r in rows]
        missing = sum(1 for v in values if v is None or str(v).strip() == "")
        miss_ratio = (missing / total) if total else 0.0
        nums = [n for n in (_to_float(v) for v in values) if n is not None]
        unique_count = len({str(v) for v in values if v is not None and str(v).strip() != ""})
        inferred = "number" if nums and len(nums) >= max(2, int(0.6 * max(total, 1))) else "string"
        sample = [str(v) for v in values if v is not None and str(v).strip() != ""][:3]
        col_stats.append(
            {
                "name": col,
                "type": inferred,
                "missing_count": missing,
                "missing_ratio": round(miss_ratio, 4),
                "unique_count": unique_count,
                "sample_values": sample,
            }
        )
        if missing == total and total > 0:
            warnings.append(f"列 `{col}` 全为空值。")
        if miss_ratio >= 0.5 and total > 0:
            warnings.append(f"列 `{col}` 缺失率较高（{miss_ratio:.1%}）。")
        if unique_count == 1 and total > 1:
            warnings.append(f"列 `{col}` 可能是常量列。")
        if unique_count == total and total > 10:
            warnings.append(f"列 `{col}` 可能是 ID 列（唯一值过多）。")

        if nums:
            numeric_summary[col] = {
                "count": len(nums),
                "min": min(nums),
                "max": max(nums),
                "mean": round(sum(nums) / len(nums), 6),
                "median": round(statistics.median(nums), 6),
                "std": round(statistics.pstdev(nums), 6) if len(nums) > 1 else 0.0,
            }
        else:
            freq = Counter(str(v) for v in values if v is not None and str(v).strip() != "").most_common(5)
            if freq:
                category_summary[col] = freq

    # 重复行检测
    sig = [json.dumps({c: r.get(c) for c in columns}, ensure_ascii=False, sort_keys=True) for r in rows]
    dup_count = max(0, len(sig) - len(set(sig)))

    return {
        "row_count": total,
        "column_count": len(columns),
        "columns": columns,
        "column_stats": col_stats,
        "numeric_summary": numeric_summary,
        "category_top_values": category_summary,
        "duplicate_rows": dup_count,
        "warnings": warnings,
    }

def _profile_to_markdown(file_info: Dict[str, Any], profile: Dict[str, Any]) -> str:
    lines: List[str] = []
    lines.append(f"# 数据概览报告（{file_info.get('original_name', '')}）")
    lines.append("")
    lines.append("## 文件基本信息")
    lines.append(f"- file_id: `{file_info.get('file_id', '')}`")
    lines.append(f"- 原始文件名: `{file_info.get('original_name', '')}`")
    lines.append(f"- 扩展名: `{file_info.get('extension', '')}`")
    lines.append(f"- 大小: `{file_info.get('size', 0)}` 字节")
    lines.append("")
    if profile.get("type") == "text":
        lines.append("## 文本统计")
        lines.append(f"- 字符数: {profile.get('char_count', 0)}")
        lines.append(f"- 行数: {profile.get('line_count', 0)}")
        preview = profile.get("preview_lines", []) or []
        lines.append("")
        lines.append("## 前几行预览")
        lines.append("```text")
        lines.extend(preview[:12] or ["(空文件)"])
        lines.append("```")
    else:
        lines.append("## 数据规模")
        lines.append(f"- 行数: {profile.get('row_count', 0)}")
        lines.append(f"- 列数: {profile.get('column_count', 0)}")
        lines.append("")
        lines.append("## 字段摘要")
        for c in profile.get("column_stats", [])[:100]:
            lines.append(
                f"- `{c.get('name')}` 类型={c.get('type')} 缺失={c.get('missing_count')} ({float(c.get('missing_ratio', 0))*100:.1f}%) 唯一值={c.get('unique_count')}"
            )
        if profile.get("numeric_summary"):
            lines.append("")
            lines.append("## 数值统计")
            for col, stat in profile.get("numeric_summary", {}).items():
                lines.append(
                    f"- `{col}`: min={stat.get('min')}, max={stat.get('max')}, mean={stat.get('mean')}, median={stat.get('median')}, std={stat.get('std')}"
                )
        if profile.get("warnings"):
            lines.append("")
            lines.append("## 异常提示")
            for w in profile.get("warnings", []):
                lines.append(f"- {w}")
    lines.append("")
    lines.append("## 后续建议问题")
    lines.append("- 哪些字段最可能与目标事件相关？")
    lines.append("- 是否需要按时间或区域进一步分组统计？")
    lines.append("- 是否需要导出清洗建议和字段映射？")
    return "\n".join(lines)

def run_data_profile(query: str, agent) -> ToolResult:
    workspace = getattr(agent, "file_workspace", None)
    if workspace is None:
        return ToolResult(success=False, content="", error="系统未启用文件工作区")

    active_file_id = getattr(agent, "active_file_id", None)
    file_info = workspace.resolve_file_from_query(query, fallback_file_id=active_file_id)
    if not file_info:
        return ToolResult(
            success=False,
            content="未找到可分析文件。请先上传文件，并在问题中提供 file_id。",
            metadata={"tool": "DATA_PROFILE"},
            error="file_not_found",
        )

    path = Path(file_info["abs_path"]).resolve()
    uploads_dir = Path(getattr(workspace, "uploads_dir", "")).resolve() if getattr(workspace, "uploads_dir", None) else None
    if uploads_dir is not None and uploads_dir not in path.parents:
        metadata = {"tool": "DATA_PROFILE", "file": {k: v for k, v in file_info.items() if k != "abs_path"}}
        return ToolResult(
            success=False,
            content="文件路径不在允许目录中。",
            metadata=metadata,
            error="invalid_file_path",
        )
    ext = file_info.get("extension", "").lower()
    metadata: Dict[str, Any] = {"tool": "DATA_PROFILE", "file": {k: v for k, v in file_info.items() if k != "abs_path"}}
    artifacts: List[Dict[str, Any]] = []
    try:
        if ext == ".csv":
            cols, rows = _read_csv(path)
            profile = _profile_table(cols, rows)
        elif ext in {".xlsx", ".xls"}:
            cols, rows, sheet = _read_xlsx(path)
            profile = _profile_table(cols, rows)
            profile["sheet"] = sheet
        elif ext == ".json":
            cols, rows, raw = _read_json(path)
            if rows:
                profile = _profile_table(cols, rows)
            else:
                profile = {
                    "type": "json",
                    "top_level_type": type(raw).__name__,
                    "keys": sorted(raw.keys()) if isinstance(raw, dict) else [],
                }
        elif ext == ".txt":
            profile = _read_txt(path)
        else:
            return ToolResult(success=False, content="", metadata=metadata, error=f"不支持的文件类型：{ext}")

        metadata["profile"] = profile
        report_md = _profile_to_markdown(file_info, profile)
        artifact = workspace.save_artifact(
            name=f"data_profile_{file_info['file_id']}.md",
            content=report_md,
            metadata={"tool": "DATA_PROFILE", "file_id": file_info["file_id"]},
        )
        artifacts.append(artifact)

        summary_lines = [
            f"[DATA_PROFILE] 已完成文件分析：{file_info['original_name']}",
            f"- file_id: {file_info['file_id']}",
        ]
        if "row_count" in profile:
            summary_lines.append(f"- 数据规模: {profile.get('row_count', 0)} 行 x {profile.get('column_count', 0)} 列")
        if profile.get("warnings"):
            summary_lines.append(f"- 异常提示: {len(profile.get('warnings', []))} 条")
            summary_lines.extend([f"  - {w}" for w in profile.get("warnings", [])[:5]])
        summary_lines.append(f"- 报告产物: {artifact.get('url')}")
        return ToolResult(success=True, content="\n".join(summary_lines), metadata=metadata, artifacts=artifacts)
    except Exception as exc:
        return ToolResult(
            success=False,
            content=f"[DATA_PROFILE] 分析失败：{exc}",
            metadata=metadata,
            artifacts=artifacts,
            error=str(exc),
        )

class DataProfileTool(BaseTool):
    name = "DATA_PROFILE"

    def run(self, tool_input: ToolInput, agent: Any) -> ToolResult:
        return run_data_profile(tool_input.task, agent)
