from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Dict, List, Optional

from tools.base import BaseTool, ToolInput, ToolResult


def _preview_csv(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        rows = []
        for i, row in enumerate(reader):
            rows.append(row)
            if i >= 5:
                break
    header = rows[0] if rows else []
    body = rows[1:] if len(rows) > 1 else []
    return {"header": header, "rows": body}


def _preview_json(path: Path) -> Dict[str, Any]:
    data = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(data, list):
        sample = data[:2]
        return {"top_level": "list", "length": len(data), "sample": sample}
    if isinstance(data, dict):
        keys = list(data.keys())[:30]
        sample = {k: data[k] for k in keys[:5]}
        return {"top_level": "dict", "keys": keys, "sample": sample}
    return {"top_level": type(data).__name__, "sample": str(data)[:300]}


def _preview_txt(path: Path) -> Dict[str, Any]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    lines = text.splitlines()
    return {"line_count": len(lines), "char_count": len(text), "preview_lines": lines[:12]}


def _preview_excel(path: Path) -> Dict[str, Any]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except Exception as exc:  # pragma: no cover
        raise RuntimeError("读取 Excel 需要 openpyxl 依赖，请先安装 openpyxl。") from exc

    wb = load_workbook(filename=str(path), data_only=True, read_only=True)
    sheet = wb[wb.sheetnames[0]]
    rows = []
    for i, row in enumerate(sheet.iter_rows(values_only=True)):
        rows.append([v for v in row])
        if i >= 5:
            break
    return {"sheet": sheet.title, "preview_rows": rows}


def run_file_inspector(query: str, agent) -> ToolResult:
    workspace = getattr(agent, "file_workspace", None)
    if workspace is None:
        return ToolResult(success=False, content="", error="系统未启用文件工作区")

    active_file_id = getattr(agent, "active_file_id", None)
    file_info = workspace.resolve_file_from_query(query, fallback_file_id=active_file_id)
    if not file_info:
        return ToolResult(
            success=False,
            content="未找到文件。请先上传文件，并在问题中提供 file_id。",
            metadata={"tool": "FILE_INSPECTOR"},
            error="file_not_found",
        )
    path = Path(file_info["abs_path"]).resolve()
    uploads_dir = Path(getattr(workspace, "uploads_dir", "")).resolve() if getattr(workspace, "uploads_dir", None) else None
    if uploads_dir is not None and uploads_dir not in path.parents:
        metadata = {
            "tool": "FILE_INSPECTOR",
            "file": {k: v for k, v in file_info.items() if k != "abs_path"},
        }
        return ToolResult(
            success=False,
            content="文件路径不在允许目录中。",
            metadata=metadata,
            error="invalid_file_path",
        )
    ext = file_info.get("extension", "").lower()
    metadata: Dict[str, Any] = {
        "tool": "FILE_INSPECTOR",
        "file": {k: v for k, v in file_info.items() if k != "abs_path"},
    }
    try:
        if ext == ".csv":
            preview = _preview_csv(path)
        elif ext == ".json":
            preview = _preview_json(path)
        elif ext == ".txt":
            preview = _preview_txt(path)
        elif ext in {".xlsx", ".xls"}:
            preview = _preview_excel(path)
        else:
            return ToolResult(success=False, content="", metadata=metadata, error=f"不支持的文件类型：{ext}")

        metadata["preview"] = preview
        content = "\n".join(
            [
                f"[FILE_INSPECTOR] 文件检查完成：{file_info['original_name']}",
                f"- file_id: {file_info['file_id']}",
                f"- 扩展名: {file_info['extension']}",
                f"- 大小: {file_info['size']} 字节",
                f"- 预览摘要: {json.dumps(preview, ensure_ascii=False)[:500]}",
            ]
        )
        return ToolResult(success=True, content=content, metadata=metadata, artifacts=[])
    except Exception as exc:
        return ToolResult(
            success=False,
            content=f"[FILE_INSPECTOR] 检查失败：{exc}",
            metadata=metadata,
            artifacts=[],
            error=str(exc),
        )


class FileInspectorTool(BaseTool):
    name = "FILE_INSPECTOR"

    def run(self, tool_input: ToolInput, agent: Any) -> ToolResult:
        return run_file_inspector(tool_input.task, agent)
